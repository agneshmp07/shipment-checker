"""Draft commercial (export) invoice from a checked booking.

Follows the UK government's export-invoice guide: invoice number and date, seller and buyer
(and consignee), related paperwork, unit price / currency / payment / discounts, weights and
number of goods, HS codes, incoterms, package marks, country of origin, transport route,
total value, additional charges, and a declaration. Everything ShipCheck knows is filled in
from the booking; what only the seller knows comes from a short form and is otherwise
marked TO COMPLETE.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Optional

import openpyxl
from fpdf import FPDF
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel

from core import fmt
from core.check import Evaluation

TODO = "TO COMPLETE"
# Road-freight incoterms (ICC Incoterms 2020) and whether the seller pays the main carriage.
INCOTERMS = {
    "EXW": ("Ex Works", False), "FCA": ("Free Carrier", False), "CPT": ("Carriage Paid To", True),
    "CIP": ("Carriage and Insurance Paid To", True), "DAP": ("Delivered at Place", True),
    "DPU": ("Delivered at Place Unloaded", True), "DDP": ("Delivered Duty Paid", True),
}
PAYMENT_TERMS = ["Bank transfer, 30 days", "Bank transfer, in advance", "Letter of credit", "Card", "Other"]


class InvoiceDetails(BaseModel):
    """What only the seller knows. Blank fields print as TO COMPLETE."""
    seller_address: str = ""
    seller_contact: str = ""
    seller_vat_eori: str = ""
    buyer_name: str = ""
    buyer_address: str = ""
    buyer_contact: str = ""
    consignee_same_as_buyer: bool = True
    consignee_name: str = ""
    consignee_address: str = ""
    po_number: str = ""
    incoterm: str = ""
    payment_terms: str = ""
    discount_pct: float = 0.0
    insurance_gbp: float = 0.0


class InvoiceLine(BaseModel):
    sku: str
    description: str
    hs_code: str
    origin: str
    quantity: float
    unit_price: float
    total: float
    cartons: float
    net_weight_kg: float


class Invoice(BaseModel):
    number: str
    issued: date
    draft_reason: str = ""                # why it can't be final yet ("" = ready)
    seller: list[str]
    buyer: list[str]
    consignee: list[str]
    references: list[tuple[str, str]]
    lines: list[InvoiceLine]
    charges: list[tuple[str, float]]      # (label, amount) after the goods subtotal
    goods_total: float
    invoice_total: float
    packages: list[tuple[str, str]]
    transport: list[tuple[str, str]]
    missing: list[str]                    # fields the seller still has to fill in


def _v(text: str, label: str, missing: list[str]) -> str:
    if text and text.strip():
        return text.strip()
    missing.append(label)
    return TODO


def build_invoice(ev: Evaluation, d: InvoiceDetails, issued: Optional[date] = None) -> Invoice:
    b, rec = ev.booking, ev.recommendation
    issued = issued or date.today()
    missing: list[str] = []
    seller = [b.seller_name or _v("", "Seller name", missing),
              _v(d.seller_address, "Seller address", missing),
              _v(d.seller_contact, "Seller contact (email / phone)", missing),
              f"VAT / EORI: {_v(d.seller_vat_eori, 'Seller VAT / EORI number', missing)}"]
    buyer = [_v(d.buyer_name, "Buyer name", missing), _v(d.buyer_address, "Buyer address", missing),
             _v(d.buyer_contact, "Buyer contact (email / phone)", missing)]
    if d.consignee_same_as_buyer:
        consignee = ["Same as buyer"]
    else:
        consignee = [_v(d.consignee_name, "Consignee name", missing),
                     _v(d.consignee_address, "Consignee address", missing)]
    incoterm = d.incoterm if d.incoterm in INCOTERMS else ""
    term_text = (f"{incoterm} {b.destination_city or ''} (Incoterms 2020)".strip() if incoterm
                 else _v("", "Incoterms agreed with the buyer", missing))
    references = [("Booking", b.booking_id or "-"),
                  ("Purchase order", _v(d.po_number, "Purchase order / proforma number", missing)),
                  ("Incoterms", term_text),
                  ("Payment terms", _v(d.payment_terms, "Payment terms", missing)),
                  ("Currency", b.currency or "GBP")]

    lines = []
    for it in b.items:
        qty, price = it.quantity or 0, it.unit_price or 0
        lines.append(InvoiceLine(
            sku=it.sku or "-", description=it.description or TODO, hs_code=it.hs_code or "MISSING",
            origin=it.country_of_origin or "MISSING", quantity=qty, unit_price=price, total=round(qty * price, 2),
            cartons=it.cartons or 0, net_weight_kg=round(qty * (it.unit_weight_kg or 0), 2)))
    goods = round(sum(l.total for l in lines), 2)

    charges: list[tuple[str, float]] = []
    if d.discount_pct:
        charges.append((f"Discount {fmt.num(d.discount_pct)}%", -round(goods * d.discount_pct / 100, 2)))
    if incoterm and INCOTERMS[incoterm][1] and rec.winner:
        charges.append((f"Freight to {b.destination_city} ({rec.winner.option.carrier}, "
                        f"{fmt.num(b.pallets)} pallet{'s' if b.pallets != 1 else ''})", rec.winner.total_cost))
    if d.insurance_gbp:
        charges.append(("Insurance", round(d.insurance_gbp, 2)))
    total = round(goods + sum(a for _, a in charges), 2)

    net = sum(l.net_weight_kg for l in lines)
    packages = [("Packages", f"{fmt.num(b.cartons)} cartons on {fmt.num(b.pallets)} pallet(s)"),
                ("Marks", f"{b.booking_id} / carton 1 to {fmt.num(b.cartons)}"),
                ("Net weight", fmt.kg(net)), ("Gross weight", fmt.kg(b.total_gross_weight_kg))]
    transport = [("From", f"{b.origin_city}, {b.origin_country}"),
                 ("To", f"{b.destination_city}, {b.destination_country} ({(b.destination_type or '').lower()})"),
                 ("Carrier", f"{rec.winner.option.carrier} ({rec.winner.option.service}), road"
                  if rec.winner else "Not yet chosen"),
                 ("Collection", fmt.d(b.ready_date))]

    reds = [i for i in ev.issues if i.severity == "Red"]
    draft = (f"{len(reds)} issue{'s' if len(reds) != 1 else ''} must be fixed first: "
             + "; ".join(sorted({i.title for i in reds})) if reds else "")
    return Invoice(number=f"INV-{b.booking_id or 'NEW'}-{issued.strftime('%Y%m%d')}", issued=issued,
                   draft_reason=draft, seller=seller, buyer=buyer, consignee=consignee, references=references,
                   lines=lines, charges=charges, goods_total=goods, invoice_total=total, packages=packages,
                   transport=transport, missing=missing)


DECLARATION = ("I declare that the information on this invoice is true and correct, and that the "
               "contents of this shipment are as stated above.")


# ---- PDF -------------------------------------------------------------------------------------

def _latin(s: str) -> str:
    """Core PDF fonts are Latin-1 only."""
    return (s.replace("—", "-").replace("–", "-").replace("→", "to").replace("’", "'")
            .encode("latin-1", "replace").decode("latin-1"))


def _cut(s, n: int) -> str:
    """Latin-1 text, shortened so one cell can't grow taller than a page."""
    s = _latin(str(s or ""))
    return s if len(s) <= n else s[:n - 3] + "..."


def invoice_pdf(inv: Invoice) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    navy, orange, grey = (15, 34, 64), (217, 88, 10), (90, 107, 132)

    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(*navy)
    pdf.cell(0, 10, "COMMERCIAL INVOICE", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(*grey)
    pdf.cell(0, 6, _latin(f"Invoice no. {inv.number}    Date of issue {fmt.d(inv.issued)}"),
             new_x="LMARGIN", new_y="NEXT")
    if inv.draft_reason or inv.missing:
        pdf.set_text_color(*orange)
        pdf.set_font("Helvetica", "B", 10)
        why = inv.draft_reason or f"{len(inv.missing)} field(s) still to complete"
        pdf.multi_cell(0, 5, _latin(f"DRAFT - {why}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    def block(title, lines, x, w):
        pdf.set_xy(x, y0)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(*orange)
        pdf.cell(w, 5, title.upper(), new_x="LEFT", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        for ln in lines:
            pdf.set_text_color(*(orange if TODO in ln else navy))
            pdf.multi_cell(w, 4.6, _cut(ln, 160), new_x="LEFT", new_y="NEXT")
        return pdf.get_y()

    y0 = pdf.get_y()
    bottoms = [block("Seller (exporter)", inv.seller, 10, 60), block("Buyer", inv.buyer, 75, 60),
               block("Consignee", inv.consignee, 140, 60)]
    pdf.set_y(max(bottoms) + 3)

    def pairs(title, rows):
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(*orange)
        pdf.cell(0, 5, title.upper(), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        for k, v in rows:
            pdf.set_text_color(*grey)
            pdf.cell(38, 5, _latin(k))
            pdf.set_text_color(*(orange if TODO in v else navy))
            pdf.multi_cell(0, 5, _cut(v, 400), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)

    pairs("References and terms", inv.references)

    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(*navy)
    with pdf.table(col_widths=(16, 44, 18, 20, 14, 18, 22, 14, 18), text_align=(
            "LEFT", "LEFT", "LEFT", "LEFT", "RIGHT", "RIGHT", "RIGHT", "RIGHT", "RIGHT"), line_height=5) as t:
        t.row(["SKU", "Description", "HS code", "Origin", "Qty", "Unit £", "Total £", "Ctns", "Net kg"])
        for l in inv.lines:
            t.row([_cut(l.sku, 30), _cut(l.description, 150), _cut(l.hs_code, 14), _cut(l.origin, 30),
                   fmt.num(l.quantity),
                   f"{l.unit_price:,.2f}", f"{l.total:,.2f}", fmt.num(l.cartons), fmt.num(l.net_weight_kg)])
    pdf.ln(2)
    pdf.set_font("Helvetica", "", 9)
    rows = [("Goods total", inv.goods_total)] + inv.charges
    for label, amount in rows:
        pdf.cell(150, 5, _latin(label), align="R")
        pdf.cell(0, 5, f"£{amount:,.2f}", align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(150, 7, "INVOICE TOTAL (GBP)", align="R")
    pdf.cell(0, 7, f"£{inv.invoice_total:,.2f}", align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pairs("Packages and marks", inv.packages)
    pairs("Transport", inv.transport)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(*navy)
    pdf.multi_cell(0, 5, _latin(DECLARATION), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)
    pdf.cell(90, 5, "Signed: ______________________________")
    pdf.cell(0, 5, f"Date: {fmt.d(inv.issued)}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 7.5)
    pdf.set_text_color(*grey)
    pdf.multi_cell(0, 4, "Draft prepared by ShipCheck from the booking. Check every detail before use; some "
                   "markets need the invoice in the buyer's language or certified by a chamber of commerce.")
    return bytes(pdf.output())


# ---- Excel ---------------------------------------------------------------------------------

def invoice_xlsx(inv: Invoice) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Commercial invoice"
    bold, todo_fill = Font(bold=True), PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    def put(row, values, b=False):
        ws.append(values)
        for c in ws[ws.max_row]:
            if b:
                c.font = bold
            if isinstance(c.value, str) and TODO in c.value:
                c.fill = todo_fill

    put(1, ["COMMERCIAL INVOICE"], b=True)
    ws["A1"].font = Font(bold=True, size=16)
    put(2, ["Invoice no.", inv.number, "Date of issue", fmt.d(inv.issued)])
    if inv.draft_reason or inv.missing:
        put(3, ["DRAFT", inv.draft_reason or f"{len(inv.missing)} field(s) still to complete"])
    ws.append([])
    for title, lines in (("Seller (exporter)", inv.seller), ("Buyer", inv.buyer), ("Consignee", inv.consignee)):
        put(0, [title], b=True)
        for ln in lines:
            put(0, ["", ln])
    ws.append([])
    put(0, ["References and terms"], b=True)
    for k, v in inv.references:
        put(0, [k, v])
    ws.append([])
    put(0, ["SKU", "Description", "HS code", "Country of origin", "Quantity", "Unit price (£)", "Total (£)",
            "Cartons", "Net weight (kg)"], b=True)
    for l in inv.lines:
        ws.append([l.sku, l.description, l.hs_code, l.origin, l.quantity, l.unit_price, l.total, l.cartons,
                   l.net_weight_kg])
        ws.cell(ws.max_row, 3).number_format = "@"
    ws.append([])
    put(0, ["", "", "", "", "", "Goods total", inv.goods_total])
    for label, amount in inv.charges:
        put(0, ["", "", "", "", "", label, amount])
    put(0, ["", "", "", "", "", "INVOICE TOTAL (GBP)", inv.invoice_total], b=True)
    for col in "FG":
        for cell in ws[col]:
            if isinstance(cell.value, float):
                cell.number_format = "£#,##0.00"
    ws.append([])
    for title, rows in (("Packages and marks", inv.packages), ("Transport", inv.transport)):
        put(0, [title], b=True)
        for k, v in rows:
            put(0, [k, v])
        ws.append([])
    put(0, [DECLARATION])
    put(0, ["Signed:", "", "Date:", fmt.d(inv.issued)])
    for col, width in zip("ABCDEFGHI", (20, 36, 12, 18, 10, 26, 14, 9, 15)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
