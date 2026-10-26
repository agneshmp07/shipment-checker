"""Write the seller's corrected bookings back into their own Excel file.

Only cells whose value changed (by "Fix all" or a fix button) are rewritten; everything
else in the workbook, including the fixed sheets and formatting, is left as uploaded.
Changed cells are highlighted with a note saying what they were, and a "Oops-Proof
changes" sheet lists every change.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from core import fmt
from core.excel_loader import BOOKING_COLUMNS, ITEM_COLUMNS, find_sheet, load_workbook, norm
from core.models import Booking

CHANGED = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
BOOKING_FIELDS = ["seller_name", "origin_city", "destination_city", "destination_country", "destination_type",
                  "ready_date", "seller_priority", "pallets", "cartons", "total_declared_value",
                  "total_gross_weight_kg", "delivery_appointment_booked"]
ITEM_FIELDS = ["description", "hs_code", "quantity", "unit_price", "units_per_carton", "cartons",
               "unit_weight_kg", "country_of_origin"]
LABELS = {"total_declared_value": "Total declared value", "total_gross_weight_kg": "Total gross weight (kg)",
          "delivery_appointment_booked": "Delivery appointment booked", "ready_date": "Ready to ship date",
          "hs_code": "HS code", "unit_weight_kg": "Unit weight (kg)", "units_per_carton": "Units per carton",
          "unit_price": "Unit price", "quantity": "Quantity (units)"}


def _columns(ws, columns: dict[str, str]) -> dict[str, int]:
    """field -> 1-based column number, from the header row (the row with 'Booking ID')."""
    for row in ws.iter_rows(min_row=1, max_row=15):
        if any(norm(c.value) == "booking id" for c in row):
            return {columns[norm(c.value)]: c.column for c in row if norm(c.value) in columns}
    return {}


def _cell_value(field: str, v):
    """Write values the way the template expects them."""
    if v is None:
        return None
    if field == "ready_date":
        return v.strftime("%d-%m-%Y") if isinstance(v, date) else v
    if field == "delivery_appointment_booked":
        return "Yes" if v else "No"
    if field == "hs_code":
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _shown(field: str, v) -> str:
    if v is None:
        return "(blank)"
    if field == "ready_date" and isinstance(v, date):
        return fmt.d(v)
    if field == "delivery_appointment_booked":
        return "Yes" if v else "No"
    if field in ("total_declared_value", "unit_price"):
        return fmt.money(v)
    if field in ("total_gross_weight_kg", "unit_weight_kg"):
        return fmt.kg(v)
    if isinstance(v, float):
        return fmt.num(v)
    return str(v)


def corrected_workbook(original: bytes, bookings: list[Booking]) -> tuple[bytes, list[dict]]:
    """Return (xlsx bytes, [change rows]) for the uploaded file with the current booking values."""
    before = load_workbook(original).bookings
    old_b = {b.source_row: b for b in before}
    old_i = {it.source_row: it for b in before for it in b.items}

    wb = openpyxl.load_workbook(BytesIO(original))
    ws_b, ws_i = find_sheet(wb, "booking", 0), find_sheet(wb, "item", 1)
    cols_b, cols_i = _columns(ws_b, BOOKING_COLUMNS), _columns(ws_i, ITEM_COLUMNS)
    changes: list[dict] = []

    def write(ws, cols, row, field, old, new, booking_id, sku, sheet):
        if field not in cols or row is None:
            return
        cell = ws.cell(row, cols[field])
        cell.value = _cell_value(field, new)
        if field == "hs_code":
            cell.number_format = "@"            # keep leading zeros as typed
        cell.fill = CHANGED
        cell.comment = Comment(f"Oops-Proof changed this from {_shown(field, old)} to {_shown(field, new)}.",
                               "Oops-Proof")
        changes.append({"Booking": booking_id, "SKU": sku or "", "Sheet": sheet, "Row": row,
                        "Field": LABELS.get(field, field.replace("_", " ").capitalize()),
                        "Was": _shown(field, old), "Now": _shown(field, new)})

    for b in bookings:
        ob = old_b.get(b.source_row)
        if ob is None:
            continue
        for f in BOOKING_FIELDS:
            if getattr(ob, f) != getattr(b, f):
                write(ws_b, cols_b, b.source_row, f, getattr(ob, f), getattr(b, f), b.booking_id, None, "Bookings")
        for it in b.items:
            oi = old_i.get(it.source_row)
            if oi is None:
                continue
            for f in ITEM_FIELDS:
                if getattr(oi, f) != getattr(it, f):
                    write(ws_i, cols_i, it.source_row, f, getattr(oi, f), getattr(it, f), b.booking_id, it.sku, "Items")

    if changes:
        log = wb.create_sheet("Oops-Proof changes")
        log.append(["Changes made by Oops-Proof — highlighted in yellow on the Bookings and Items sheets"])
        log["A1"].font = Font(bold=True)
        headers = list(changes[0])
        log.append(headers)
        for c in log[2]:
            c.font = Font(bold=True)
        for ch in changes:
            log.append([ch[h] for h in headers])
        for col, width in zip("ABCDEFG", (12, 12, 10, 6, 26, 16, 16)):
            log.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), changes
