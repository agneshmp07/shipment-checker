"""Read Sheets 1 (Bookings) and 2 (Items) into Booking objects.

Never raises on bad data: blanks, text in number columns, unknown booking IDs
and bad dates all become Issues (rule_id "DATA"). Only a file that is not an
Excel workbook at all, or has no Bookings sheet, produces a file-level issue
and an empty booking list.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any, Optional

import openpyxl

from core.models import (
    DESTINATIONS, DESTINATION_TYPES, ITEM_ORIGIN_COUNTRIES, ORIGIN_CITIES,
    ORIGIN_COUNTRY, PRIORITIES, Booking, Issue, Item,
)

# ---- Column matching -------------------------------------------------------


def norm(text: Any) -> str:
    """'Total gross weight (kg) ' -> 'total gross weight kg'."""
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


# normalised header -> Booking field
BOOKING_COLUMNS = {
    "booking id": "booking_id",
    "seller name": "seller_name", "seller": "seller_name",
    "origin city": "origin_city",
    "origin country": "origin_country",
    "destination city": "destination_city",
    "destination country": "destination_country",
    "destination type": "destination_type",
    "ready to ship date": "ready_date", "ready date": "ready_date", "ready to ship": "ready_date",
    "seller priority": "seller_priority", "priority": "seller_priority",
    "pallets": "pallets", "number of pallets": "pallets",
    "cartons": "cartons", "total cartons": "cartons",
    "total declared value": "total_declared_value", "declared value": "total_declared_value",
    "currency": "currency",
    "total gross weight kg": "total_gross_weight_kg", "total gross weight": "total_gross_weight_kg",
    "gross weight kg": "total_gross_weight_kg",
    "delivery appointment booked": "delivery_appointment_booked",
    "delivery appointment": "delivery_appointment_booked",
}

# normalised header -> Item field
ITEM_COLUMNS = {
    "booking id": "booking_id",
    "sku": "sku",
    "description": "description",
    "hs code": "hs_code",
    "quantity units": "quantity", "quantity": "quantity",
    "unit price": "unit_price", "unit price gbp": "unit_price",
    "units per carton": "units_per_carton",
    "cartons": "cartons",
    "unit weight kg": "unit_weight_kg", "unit weight": "unit_weight_kg",
    "country of origin": "country_of_origin",
}

BOOKING_NUMBERS = {
    "pallets": "Pallets", "cartons": "Cartons",
    "total_declared_value": "Total declared value",
    "total_gross_weight_kg": "Total gross weight (kg)",
}
ITEM_NUMBERS = {
    "quantity": "Quantity", "unit_price": "Unit price",
    "units_per_carton": "Units per carton", "cartons": "Cartons",
    "unit_weight_kg": "Unit weight (kg)",
}

# ---- Value parsing ---------------------------------------------------------


def is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def parse_number(v: Any) -> tuple[Optional[float], Optional[str]]:
    """Return (number, error). Accepts 5400, '5,400', '£5,400.00', '620 kg'."""
    if is_blank(v):
        return None, "blank"
    if isinstance(v, bool):
        return None, f"'{v}' is not a number"
    if isinstance(v, (int, float)):
        n = float(v)
    else:
        s = re.sub(r"(?i)[£,\s]|gbp|kgs?$", "", str(v).strip())
        try:
            n = float(s)
        except ValueError:
            return None, f"'{v}' is not a number"
    if not math.isfinite(n):                     # 'nan', 'inf', '1e999'
        return None, f"'{v}' is not a number"
    return n, None


def parse_date(v: Any) -> tuple[Optional[date], Optional[str]]:
    """Excel datetime, Excel serial number, or text DD-MM-YYYY (also / or .)."""
    if is_blank(v):
        return None, "blank"
    if isinstance(v, datetime):
        return v.date(), None
    if isinstance(v, date):
        return v, None
    if isinstance(v, (int, float)) and 20000 < v < 80000:   # Excel serial date
        return date(1899, 12, 30) + timedelta(days=int(v)), None
    s = str(v).strip()
    for f in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y",
              "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, f).date(), None
        except ValueError:
            pass
    return None, f"'{v}' is not a date in DD-MM-YYYY format"


def parse_hs(v: Any) -> tuple[Optional[str], bool]:
    """HS code as text. Returns (code, possible_leading_zero_lost).

    Excel often stores 090111 as the number 90111 (or 90111.0). We convert it
    back to text and flag it when the digit count (5 or 9) suggests a lost 0.
    """
    if is_blank(v):
        return None, False
    from_number = False
    if isinstance(v, float) and v.is_integer():
        s, from_number = str(int(v)), True
    elif isinstance(v, int) and not isinstance(v, bool):
        s, from_number = str(v), True
    else:
        s = str(v).strip()
        if re.fullmatch(r"\d+\.0+", s):           # '90111.0' from a CSV export
            s, from_number = s.split(".")[0], True
    digits = re.sub(r"\D", "", s)
    return s, from_number and len(digits) in (5, 9)


def parse_yes_no(v: Any) -> tuple[Optional[bool], Optional[str]]:
    if is_blank(v):
        return None, "blank"
    if isinstance(v, bool):
        return v, None
    s = str(v).strip().lower()
    if s in ("yes", "y", "true", "1"):
        return True, None
    if s in ("no", "n", "false", "0"):
        return False, None
    return None, f"'{v}' should be Yes or No"


def match_choice(v: Any, allowed: list[str]) -> Optional[str]:
    """Case/space-insensitive match against an allowed list, else None."""
    key = norm(v)
    for a in allowed:
        if norm(a) == key:
            return a
    return None


def text(v: Any) -> Optional[str]:
    if is_blank(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ---- Sheet reading ---------------------------------------------------------


def find_sheet(wb, keyword: str, position: Optional[int]):
    """Find a sheet by name keyword first, then by position."""
    for ws in wb.worksheets:
        if keyword in norm(ws.title):
            return ws
    if position is not None and position < len(wb.worksheets):
        return wb.worksheets[position]
    return None


MAX_BLANK_RUN = 1000     # stop after this many empty rows in a row (a stray cell far down the sheet)


def read_table(ws, columns: dict[str, str], notes: Optional[list[str]] = None
               ) -> tuple[list[tuple[int, dict]], list[str]]:
    """Return ([(excel_row_number, {field: raw_value})], unknown_headers).

    The header row is the first row (within the top 15) that has a
    'Booking ID' cell, so title rows above the table are tolerated.
    """
    rows, blank_run = [], 0
    for row in ws.iter_rows(values_only=True):
        rows.append(row)
        blank_run = blank_run + 1 if all(is_blank(c) for c in row) else 0
        if blank_run > MAX_BLANK_RUN:
            if notes is not None and ws.max_row > len(rows):
                notes.append(f"Stopped reading sheet “{ws.title}” after {MAX_BLANK_RUN:,} empty rows (at row "
                             f"{len(rows):,}); anything further down, up to row {ws.max_row:,}, was ignored.")
            break
    header_idx = None
    for i, row in enumerate(rows[:15]):
        if any(norm(c) == "booking id" for c in row):
            header_idx = i
            break
    if header_idx is None:
        return [], []

    col_map: dict[int, str] = {}
    unknown = []
    for j, h in enumerate(rows[header_idx]):
        key = norm(h)
        if not key:
            continue
        if key in columns:
            col_map[j] = columns[key]
        else:
            unknown.append(str(h))

    out = []
    for i in range(header_idx + 1, len(rows)):
        row = rows[i]
        rec = {f: row[j] if j < len(row) else None for j, f in col_map.items()}
        if all(is_blank(v) for v in rec.values()):
            continue
        out.append((i + 1, rec))     # +1: Excel rows are 1-based
    return out, unknown


def data_issue(severity: str, title: str, explanation: str, fix: str,
               booking_id: Optional[str] = None, sku: Optional[str] = None) -> Issue:
    return Issue(rule_id="DATA", severity=severity, title=title, explanation=explanation,
                 why_it_matters="This value is needed to check the booking; until it is "
                                "corrected some checks cannot run.",
                 fix=fix, booking_id=booking_id, item_sku=sku)


def build_booking(row_no: int, rec: dict) -> Booking:
    b = Booking(source_row=row_no, booking_id=text(rec.get("booking_id")),
                seller_name=text(rec.get("seller_name")))
    bid = b.booking_id
    issues = b.data_issues
    where = f"Bookings sheet row {row_no}"

    # Numbers
    for f, label in BOOKING_NUMBERS.items():
        val, err = parse_number(rec.get(f))
        if err:
            issues.append(data_issue(
                "Red", f"{label} is missing" if err == "blank" else f"{label} is not a number",
                f"{where}: {label} {'is blank' if err == 'blank' else '— ' + err}.",
                f"Enter {label} as a plain number.", bid))
        elif val < 0:
            issues.append(data_issue("Red", f"{label} is negative",
                                     f"{where}: {label} is {val:g}, which cannot be negative.",
                                     f"Enter the correct {label}.", bid))
        else:
            setattr(b, f, val)

    # Date
    val, err = parse_date(rec.get("ready_date"))
    if err:
        issues.append(data_issue("Red", "Ready to ship date is missing or unreadable",
                                 f"{where}: ready to ship date is {err}.",
                                 "Enter the date as DD-MM-YYYY, e.g. 05-11-2026.", bid))
    b.ready_date = val

    # Fixed lists
    raw_oc = rec.get("origin_country")
    b.origin_country = match_choice(raw_oc, [ORIGIN_COUNTRY]) if not is_blank(raw_oc) else ORIGIN_COUNTRY
    if b.origin_country is None:
        b.origin_country = text(raw_oc)
        issues.append(data_issue("Red", "Origin country not supported",
                                 f"{where}: origin country is '{raw_oc}'. Only United Kingdom is supported.",
                                 "Set origin country to United Kingdom.", bid))

    _choice(b, rec, "origin_city", "Origin city", ORIGIN_CITIES, where)
    _choice(b, rec, "destination_country", "Destination country", list(DESTINATIONS), where)
    allowed_cities = DESTINATIONS.get(b.destination_country or "",
                                      [c for cs in DESTINATIONS.values() for c in cs])
    _choice(b, rec, "destination_city", "Destination city", allowed_cities, where)
    _choice(b, rec, "destination_type", "Destination type", DESTINATION_TYPES, where)
    _choice(b, rec, "seller_priority", "Seller priority", PRIORITIES, where)

    cur = rec.get("currency")
    b.currency = "GBP" if is_blank(cur) or norm(cur) == "gbp" else text(cur)
    if b.currency != "GBP":
        issues.append(data_issue("Red", "Currency must be GBP",
                                 f"{where}: currency is '{cur}'. All values must be in GBP.",
                                 "Convert values to GBP and set currency to GBP.", bid))

    val, err = parse_yes_no(rec.get("delivery_appointment_booked"))
    if err:
        issues.append(data_issue("Amber", "Delivery appointment booked is unclear",
                                 f"{where}: delivery appointment booked is {err}.",
                                 "Enter Yes or No.", bid))
    b.delivery_appointment_booked = val

    if not bid:
        issues.append(data_issue("Red", "Booking ID is missing",
                                 f"{where} has no Booking ID, so its items cannot be matched.",
                                 "Give the booking an ID (e.g. BK-001) and use it on its item rows."))
    if not b.seller_name:
        issues.append(data_issue("Amber", "Seller name is missing", f"{where}: seller name is blank.",
                                 "Enter the seller's business name.", bid))
    return b


def _choice(b: Booking, rec: dict, f: str, label: str, allowed: list[str], where: str):
    raw = rec.get(f)
    if is_blank(raw):
        b.data_issues.append(data_issue("Red", f"{label} is missing", f"{where}: {label} is blank.",
                                        f"Choose one of: {', '.join(allowed)}.", b.booking_id))
        return
    match = match_choice(raw, allowed)
    if match:
        setattr(b, f, match)
    else:
        setattr(b, f, text(raw))
        b.data_issues.append(data_issue("Red", f"{label} not recognised",
                                        f"{where}: {label} '{raw}' is not one of the allowed values.",
                                        f"Choose one of: {', '.join(allowed)}.", b.booking_id))


def build_item(row_no: int, rec: dict) -> tuple[Item, list[Issue]]:
    hs, lost_zero = parse_hs(rec.get("hs_code"))
    it = Item(source_row=row_no, booking_id=text(rec.get("booking_id")), sku=text(rec.get("sku")),
              description=text(rec.get("description")), hs_code=hs,
              hs_possible_leading_zero_lost=lost_zero)
    issues: list[Issue] = []
    label_sku = it.sku or f"row {row_no}"
    where = f"Items sheet row {row_no} ({label_sku})"

    for f, label in ITEM_NUMBERS.items():
        val, err = parse_number(rec.get(f))
        if err:
            issues.append(data_issue(
                "Red", f"{label} is missing for {label_sku}" if err == "blank"
                else f"{label} is not a number for {label_sku}",
                f"{where}: {label} {'is blank' if err == 'blank' else '— ' + err}.",
                f"Enter {label} as a plain number.", it.booking_id, it.sku))
        elif val < 0:
            issues.append(data_issue("Red", f"{label} is negative for {label_sku}",
                                     f"{where}: {label} is {val:g}, which cannot be negative.",
                                     f"Enter the correct {label}.", it.booking_id, it.sku))
        else:
            setattr(it, f, val)

    # Country of origin: blank is R-09's job; an unknown value is a data issue.
    raw_coo = rec.get("country_of_origin")
    if not is_blank(raw_coo):
        it.country_of_origin = match_choice(raw_coo, ITEM_ORIGIN_COUNTRIES) or text(raw_coo)
        if it.country_of_origin not in ITEM_ORIGIN_COUNTRIES:
            issues.append(data_issue("Amber", f"Country of origin not recognised for {label_sku}",
                                     f"{where}: country of origin '{raw_coo}' is not in the allowed list.",
                                     f"Choose one of: {', '.join(ITEM_ORIGIN_COUNTRIES)}.",
                                     it.booking_id, it.sku))
    if not it.sku:
        issues.append(data_issue("Amber", f"SKU is missing on Items row {row_no}",
                                 f"{where}: SKU is blank.", "Enter the product code.", it.booking_id))
    if not it.description:
        issues.append(data_issue("Amber", f"Description is missing for {label_sku}",
                                 f"{where}: description is blank, so customs cannot tell what it is.",
                                 "Enter a plain description, e.g. 'Ceramic coffee mug, 350 ml'.",
                                 it.booking_id, it.sku))
    return it, issues


# ---- Public API ------------------------------------------------------------


@dataclass
class LoadResult:
    bookings: list[Booking] = field(default_factory=list)
    file_issues: list[Issue] = field(default_factory=list)   # not tied to one booking
    notes: list[str] = field(default_factory=list)            # informational, not problems


# The template's grey example rows (row 3 of Bookings and Items). Skipped only when a
# row matches them exactly, so real data typed over row 3 is still checked.
EXAMPLE_BOOKING = dict(booking_id="BK-001", seller_name="Northfield Goods", origin_city="Manchester",
                       destination_city="Frankfurt", destination_type="Warehouse",
                       ready_date=date(2026, 11, 5), seller_priority="Cheapest", pallets=3, cartons=90,
                       total_declared_value=5400, total_gross_weight_kg=620, delivery_appointment_booked=True)
EXAMPLE_ITEM = dict(booking_id="BK-001", sku="MUG-01", description="Ceramic coffee mug, 350 ml",
                    hs_code="691200", quantity=1080, unit_price=5, units_per_carton=12, cartons=90,
                    unit_weight_kg=0.5, country_of_origin="United Kingdom")


def _matches(obj, example: dict) -> bool:
    return all(getattr(obj, k) == v for k, v in example.items())


def load_workbook(source) -> LoadResult:
    """source: a path, bytes, or a file-like object (e.g. Streamlit upload)."""
    result = LoadResult()
    try:
        if isinstance(source, (bytes, bytearray)):
            source = BytesIO(source)
        wb = openpyxl.load_workbook(source, data_only=True)
    except Exception as e:   # noqa: BLE001 - any failure here means "not a usable workbook"
        result.file_issues.append(data_issue("Red", "File could not be opened",
                                             f"The file could not be read as an Excel workbook ({e}).",
                                             "Upload the .xlsx template with the Bookings and Items sheets."))
        return result

    ws_b = find_sheet(wb, "booking", 0)
    ws_i = find_sheet(wb, "item", 1)
    b_rows, b_unknown = read_table(ws_b, BOOKING_COLUMNS, result.notes) if ws_b else ([], [])
    i_rows, i_unknown = (read_table(ws_i, ITEM_COLUMNS, result.notes) if ws_i and ws_i is not ws_b
                         else ([], []))

    if not b_rows:
        result.file_issues.append(data_issue("Red", "No bookings found",
                                             "Could not find a Bookings sheet with a 'Booking ID' column and data rows.",
                                             "Use the template's Bookings sheet (Sheet 1) without renaming its columns."))
        return result
    for label, unknown in (("Bookings", b_unknown), ("Items", i_unknown)):
        if unknown:
            result.file_issues.append(data_issue("Amber", f"Unrecognised columns in {label} sheet",
                                                 f"These columns were ignored: {', '.join(unknown)}.",
                                                 "Check the column names match the template."))

    by_id: dict[str, Booking] = {}
    skipped_example = False
    for row_no, rec in b_rows:
        b = build_booking(row_no, rec)
        if _matches(b, EXAMPLE_BOOKING):
            skipped_example = True
            result.notes.append(f"Skipped Bookings row {row_no}: it is the template's grey example row (BK-001).")
            continue
        if b.booking_id and b.booking_id in by_id:
            result.file_issues.append(data_issue(
                "Red", f"Duplicate Booking ID {b.booking_id}",
                f"Booking ID {b.booking_id} appears on Bookings rows {by_id[b.booking_id].source_row} "
                f"and {row_no}. Only the first is checked.",
                "Give each booking a unique ID.", b.booking_id))
            continue
        if b.booking_id:
            by_id[b.booking_id] = b
        result.bookings.append(b)

    for row_no, rec in i_rows:
        item, issues = build_item(row_no, rec)
        if skipped_example and "BK-001" not in by_id and _matches(item, EXAMPLE_ITEM):
            result.notes.append(f"Skipped Items row {row_no}: it is the template's grey example row (MUG-01).")
            continue
        owner = by_id.get(item.booking_id or "")
        if owner is None:
            result.file_issues.append(data_issue(
                "Red", f"Item {item.sku or f'on row {row_no}'} has unknown Booking ID",
                f"Items sheet row {row_no}: Booking ID '{item.booking_id or '(blank)'}' does not "
                f"match any booking in the Bookings sheet ({', '.join(by_id) or 'none'}).",
                "Correct the Booking ID on this item row.", item.booking_id, item.sku))
            continue
        owner.items.append(item)
        owner.data_issues.extend(issues)

    if not result.bookings and skipped_example:
        result.file_issues.append(data_issue("Red", "Only the example row was found",
                                             "The Bookings sheet contains only the template's grey example row (BK-001).",
                                             "Enter your bookings from row 4 of the Bookings sheet."))
    for b in result.bookings:
        if not b.items:
            b.data_issues.append(data_issue("Red", "Booking has no items",
                                            f"No rows in the Items sheet use Booking ID {b.booking_id}.",
                                            "Add the products in this shipment to the Items sheet.",
                                            b.booking_id))
    return result


def dump_workbook(path: str, max_rows: int = 8) -> None:
    """Print every sheet's name, header row and first rows (for inspection)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print(f"\n=== Sheet '{ws.title}' ({ws.max_row} rows x {ws.max_column} cols) ===")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            print("  ", [c for c in row])


if __name__ == "__main__":
    import sys

    from core import fmt
    sys.stdout.reconfigure(encoding="utf-8")
    path = sys.argv[1]
    dump_workbook(path)
    res = load_workbook(path)
    print(f"\n=== Parsed: {len(res.bookings)} bookings ===")
    for b in res.bookings:
        print(f"\n{b.booking_id} | {b.seller_name} | {b.route} | {b.destination_type} | "
              f"ready {fmt.d(b.ready_date)} | {b.seller_priority} | pallets {fmt.num(b.pallets)} | "
              f"cartons {fmt.num(b.cartons)} | {fmt.money(b.total_declared_value)} | "
              f"{fmt.kg(b.total_gross_weight_kg)} | appt {b.delivery_appointment_booked}")
        for it in b.items:
            print(f"   - {it.sku} | {it.description} | HS {it.hs_code}"
                  f"{' (leading 0 lost?)' if it.hs_possible_leading_zero_lost else ''} | qty {fmt.num(it.quantity)}"
                  f" x {fmt.money(it.unit_price)} | {fmt.num(it.units_per_carton)}/ctn x {fmt.num(it.cartons)} ctn"
                  f" | {fmt.kg(it.unit_weight_kg)} | {it.country_of_origin}")
        for iss in b.data_issues:
            print(f"   ! [{iss.severity}] {iss.title}: {iss.explanation}")
    for iss in res.file_issues:
        print(f"! FILE [{iss.severity}] {iss.title}: {iss.explanation}")
