"""Loader tests using workbooks built in the template's column format."""
from datetime import date, datetime

import openpyxl
import pytest

from core.excel_loader import load_workbook, parse_hs

BOOKING_HEADERS = ["Booking ID", "Seller name", "Origin city", "Origin country", "Destination city",
                   "Destination country", "Destination type", "Ready to ship date", "Seller priority",
                   "Pallets", "Cartons", "Total declared value", "Currency", "Total gross weight (kg)",
                   "Delivery appointment booked"]
ITEM_HEADERS = ["Booking ID", "SKU", "Description", "HS code", "Quantity (units)", "Unit price",
                "Units per carton", "Cartons", "Unit weight (kg)", "Country of origin"]

GOOD_BOOKING = ["BK-001", "Northfield Goods", "Manchester", "United Kingdom", "Frankfurt", "Germany",
                "Warehouse", "05-11-2026", "Cheapest", 3, 90, 5400, "GBP", 620, "Yes"]
GOOD_ITEM = ["BK-001", "MUG-01", "Ceramic coffee mug, 350 ml", 691200, 1080, 5.00, 12, 90, 0.5,
             "United Kingdom"]


def make_xlsx(tmp_path, bookings, items, booking_headers=BOOKING_HEADERS, title_row=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bookings"
    if title_row:
        ws.append(["Sheet 1 - Bookings"])
    ws.append(booking_headers)
    for r in bookings:
        ws.append(r)
    wi = wb.create_sheet("Items")
    wi.append(ITEM_HEADERS)
    for r in items:
        wi.append(r)
    p = tmp_path / "t.xlsx"
    wb.save(p)
    return p


def test_brief_example_parses(tmp_path):
    res = load_workbook(make_xlsx(tmp_path, [GOOD_BOOKING], [GOOD_ITEM]))
    assert res.file_issues == []
    (b,) = res.bookings
    assert b.data_issues == []
    assert b.ready_date == date(2026, 11, 5)
    assert b.pallets == 3 and b.total_declared_value == 5400 and b.delivery_appointment_booked is True
    (it,) = b.items
    assert it.hs_code == "691200" and not it.hs_possible_leading_zero_lost
    assert it.quantity == 1080 and it.cartons == 90


def test_headers_case_and_spaces_and_title_row(tmp_path):
    messy = [f"  {h.upper()}  " for h in BOOKING_HEADERS]
    res = load_workbook(make_xlsx(tmp_path, [GOOD_BOOKING], [GOOD_ITEM], messy, title_row=True))
    assert len(res.bookings) == 1 and res.bookings[0].seller_priority == "Cheapest"


def test_excel_datetime_and_lowercase_choices(tmp_path):
    row = GOOD_BOOKING.copy()
    row[7] = datetime(2026, 11, 5)
    row[6], row[8], row[14] = "warehouse", "most reliable", "no"
    b = load_workbook(make_xlsx(tmp_path, [row], [GOOD_ITEM])).bookings[0]
    assert b.ready_date == date(2026, 11, 5)
    assert b.destination_type == "Warehouse" and b.seller_priority == "Most reliable"
    assert b.delivery_appointment_booked is False


def test_hs_leading_zero_lost():
    assert parse_hs(90111) == ("90111", True)
    assert parse_hs(90111.0) == ("90111", True)
    assert parse_hs("090111") == ("090111", False)
    assert parse_hs(691200) == ("691200", False)
    assert parse_hs(None) == (None, False)


def test_bad_values_become_issues_not_crashes(tmp_path):
    row = GOOD_BOOKING.copy()
    row[9], row[7], row[11] = "three", "31-02-2026", None
    item = GOOD_ITEM.copy()
    item[4] = "lots"
    b = load_workbook(make_xlsx(tmp_path, [row], [item])).bookings[0]
    titles = " | ".join(i.title for i in b.data_issues)
    assert "Pallets is not a number" in titles
    assert "Ready to ship date" in titles
    assert "Total declared value is missing" in titles
    assert "Quantity is not a number for MUG-01" in titles
    assert b.pallets is None and b.items[0].quantity is None


def test_unknown_booking_id_and_booking_without_items(tmp_path):
    second = GOOD_BOOKING.copy()
    second[0] = "BK-002"
    orphan = GOOD_ITEM.copy()
    orphan[0] = "BK-999"
    res = load_workbook(make_xlsx(tmp_path, [GOOD_BOOKING, second], [GOOD_ITEM, orphan]))
    assert any("unknown Booking ID" in i.title and "BK-999" in i.explanation for i in res.file_issues)
    bk2 = next(b for b in res.bookings if b.booking_id == "BK-002")
    assert any(i.title == "Booking has no items" for i in bk2.data_issues)


def test_blank_hs_and_country_are_left_for_rules(tmp_path):
    item = GOOD_ITEM.copy()
    item[3], item[9] = None, None
    b = load_workbook(make_xlsx(tmp_path, [GOOD_BOOKING], [item])).bookings[0]
    assert b.items[0].hs_code is None and b.items[0].country_of_origin is None
    assert b.data_issues == []        # R-01 / R-09 report these, not the loader


def test_not_an_excel_file():
    res = load_workbook(b"this is not a workbook")
    assert res.bookings == [] and res.file_issues[0].title == "File could not be opened"
