"""Phase 9: carrier import, no-crash guarantees, and the no-API-key fallback."""
import json
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from core.check import evaluate
from core.excel_loader import load_workbook
from scripts.build_carriers import parse
from tests.conftest import STANDIN_CARRIERS
from tests.fixtures.make_planted_mistakes import BOOKING_HEADERS, ITEM_HEADERS, OUT, build

ROOT = Path(__file__).resolve().parent.parent
TODAY = date(2026, 10, 25)


# ---- Sheet 3 -> carriers.json ---------------------------------------------------

def _carrier_book(tmp_path, header, rows, title="Carrier options"):
    wb = openpyxl.Workbook()
    wb.active.title = "Bookings"
    wb.create_sheet("Items")
    ws = wb.create_sheet(title)
    ws.append(["Sheet 3 - Carrier options"])
    ws.append(header)
    for r in rows:
        ws.append(r)
    p = tmp_path / "template.xlsx"
    wb.save(p)
    return str(p)


def test_build_carriers_separate_columns(tmp_path):
    p = _carrier_book(tmp_path, ["Carrier", "Service", "Origin country", "Destination country",
                                 "Price per pallet", "Transit time (days)", "On-time rate (%)",
                                 "Max weight per pallet (kg)"],
                      [["Carrier A", "Road economy", "United Kingdom", "Germany", 180, 5, 92, 500]])
    carriers, problems = parse(p)
    assert problems == []
    assert carriers[0].model_dump() == {"carrier": "Carrier A", "service": "Road economy",
                                        "origin_country": "United Kingdom", "destination_country": "Germany",
                                        "price_per_pallet": 180, "transit_days": 5, "on_time_rate": 92,
                                        "max_weight_per_pallet_kg": 500}


def test_build_carriers_combined_columns_and_fraction_rate(tmp_path):
    p = _carrier_book(tmp_path, ["Carrier · Service", "Origin country → Destination country", "Price per pallet",
                                 "Transit time (days)", "On-time rate (%)", "Max weight per pallet (kg)"],
                      [["Carrier B · Express", "United Kingdom → France", "£250", 3, 0.96, 400]])
    (c,), problems = parse(p)
    assert problems == [] and (c.carrier, c.service, c.destination_country, c.on_time_rate) == \
        ("Carrier B", "Express", "France", 96)


def test_build_carriers_refuses_incomplete_rows(tmp_path):
    p = _carrier_book(tmp_path, ["Carrier", "Service", "Origin country", "Destination country",
                                 "Price per pallet", "Transit time (days)", "On-time rate (%)",
                                 "Max weight per pallet (kg)"],
                      [["Carrier A", "Road", "United Kingdom", "Germany", "tbc", 5, 92, 500]])
    _, problems = parse(p)
    assert problems and "price_per_pallet" in problems[0]


# ---- Never crash on bad input --------------------------------------------------------

JUNK_BOOKINGS = [
    ["BK-J1", None, "Glasgow", "Scotland", "Madrid", "Spain", "Garage", "yesterday", "Whatever",
     0, -5, "lots", "EUR", "", "Maybe"],
    ["BK-J2", "Zero Pallets", "London", "United Kingdom", "Paris", "France", "Warehouse", "31-02-2026",
     "Fastest", 0, 10, 100, "GBP", 50, "No"],
    [None, "No ID", "Leeds", "United Kingdom", "Lyon", "France", "Business address", datetime(2026, 12, 1),
     "Cheapest", 1, 1, 10, "GBP", 5, "Yes"],
    ["BK-J4", "Huge", "Leeds", "United Kingdom", "Lyon", "France", "Business address", 46000,
     "Balanced", 1, 1, 1e12, "GBP", 1e9, "Yes"],
    ["BK-J5", "No items", "Leeds", "United Kingdom", "Lyon", "France", "Business address", "01-12-2026",
     "Balanced", 1, 1, 10, "GBP", 5, "Yes"],
]
JUNK_ITEMS = [
    ["BK-J1", None, None, "abc", "ten", None, 0, 0, None, "Mars"],
    ["BK-J2", "X", "Lithium battery pack", 85076000, 10, 1, 0, 1, 0.5, "China"],
    ["BK-J4", "Y", "Thing", 1.5, 1e6, 1e6, 1, 1e6, 1000, "India"],
    ["BK-NOPE", "Z", "Orphan", 123456, 1, 1, 1, 1, 1, "China"],
    [None, None, None, None, None, None, None, None, None, None],
]


def test_junk_workbook_never_crashes(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bookings"
    ws.append(BOOKING_HEADERS + ["Unexpected column"])
    for r in JUNK_BOOKINGS:
        ws.append(r)
    wi = wb.create_sheet("Items")
    wi.append(ITEM_HEADERS)
    for r in JUNK_ITEMS:
        wi.append(r)
    p = tmp_path / "junk.xlsx"
    wb.save(p)
    res = load_workbook(p)
    assert len(res.bookings) == 5 and res.file_issues
    for b in res.bookings:
        ev = evaluate(b, STANDIN_CARRIERS, TODAY)
        assert ev.status == "Blocked" or b.booking_id in ("BK-J5",)   # every junk booking is caught
        assert ev.summary and ev.next_steps


@pytest.fixture
def carriers_env(tmp_path, monkeypatch):
    path = tmp_path / "carriers.json"
    path.write_text(json.dumps([c.model_dump() for c in STANDIN_CARRIERS]))
    monkeypatch.setenv("SHIPMENT_CARRIERS_PATH", str(path))
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)    # also proves the no-key path


def _render_all_script(path):
    import streamlit as st
    from datetime import date
    from core.carriers import load_carriers
    from core.check import evaluate
    from core.excel_loader import load_workbook
    from ui.results import render_evaluation
    for n, b in enumerate(load_workbook(path).bookings):
        render_evaluation(evaluate(b, load_carriers(), date.today()), key=f"b{n}")


def test_every_fixture_booking_renders_without_key(carriers_env):
    build()
    at = AppTest.from_function(_render_all_script, args=(str(OUT),), default_timeout=60)
    at.run()
    assert not at.exception
    assert len([s for s in at.subheader if s.value.startswith("BK-")]) == 14


def test_every_junk_booking_renders(carriers_env, tmp_path):
    test_junk_workbook_never_crashes(tmp_path)
    at = AppTest.from_function(_render_all_script, args=(str(tmp_path / "junk.xlsx"),), default_timeout=60)
    at.run()
    assert not at.exception


def test_upload_and_chat_work_without_key_or_carriers(monkeypatch, tmp_path):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setenv("SHIPMENT_CARRIERS_PATH", str(tmp_path / "missing.json"))
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    at.switch_page("app_pages/upload_file.py").run()
    at.button[0].click().run()
    assert not at.exception
    assert any("rate card" in w.value for w in at.warning)        # clear message, no crash
    at.switch_page("app_pages/chat_booking.py").run()
    at.chat_input[0].set_value("1080 mugs from Manchester to a warehouse in Frankfurt").run()
    assert not at.exception
