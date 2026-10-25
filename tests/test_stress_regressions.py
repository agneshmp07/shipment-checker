"""Bugs found by the stress tests: each one hung or crashed the app before it was fixed."""
import io
import time
from datetime import date

import openpyxl

from core.chat_agent import Extraction, ItemPatch
from core.check import evaluate
from core.diagnosis import fix_all
from core.excel_loader import MAX_BLANK_RUN, load_workbook, parse_number
from core.invoice import InvoiceDetails, build_invoice, invoice_pdf
from core.labels import MAX_LABELS, carton_list, label_count, labels_pdf
from tests.conftest import STANDIN_CARRIERS
from tests.test_rules import booking, item


def test_absurd_carton_count_does_not_hang_labels():
    b = booking([item(cartons=10**9, quantity=12 * 10**9)], cartons=10**9)
    t = time.perf_counter()
    assert label_count(b) > MAX_LABELS and carton_list(b) == []
    assert labels_pdf(b).startswith(b"%PDF") and time.perf_counter() - t < 2


def test_non_finite_cartons_make_no_labels():
    assert label_count(booking([item(cartons=float("inf")), item(sku="B", cartons=float("nan"))])) == 0


def test_nan_and_infinity_from_the_ai_are_ignored():
    ex = Extraction(pallets=float("nan"), total_gross_weight_kg=float("inf"),
                    items=[ItemPatch(quantity=float("-inf"), cartons=3)])
    assert ex.pallets is None and ex.total_gross_weight_kg is None
    assert ex.items[0].quantity is None and ex.items[0].cartons == 3


def test_nan_in_a_spreadsheet_is_not_a_number():
    for v in ("nan", float("inf"), "1e999"):
        assert parse_number(v)[0] is None


def test_very_long_text_still_makes_an_invoice():
    b = booking([item(description="x" * 2000, sku="S" * 500, country_of_origin="y" * 900)],
                seller_name="z" * 3000)
    inv = build_invoice(evaluate(b, STANDIN_CARRIERS, date.today()), InvoiceDetails(buyer_name="q" * 5000))
    assert invoice_pdf(inv).startswith(b"%PDF")


def test_fix_all_is_fast_with_hundreds_of_items():
    b = booking([item(sku=f"S{i}") for i in range(200)], total_gross_weight_kg=5)   # R-07 fails
    t = time.perf_counter()
    fix_all([b], lambda x: evaluate(x, STANDIN_CARRIERS, date.today()))
    assert time.perf_counter() - t < 5


def test_stray_cell_far_down_the_sheet_is_skipped_quickly():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bookings"
    ws.append(["Booking ID", "Seller name"])
    ws.append(["BK-1", "Acme"])
    ws.cell(500_000, 1).value = "stray"
    wb.create_sheet("Items").append(["Booking ID", "SKU"])
    out = io.BytesIO()
    wb.save(out)
    t = time.perf_counter()
    res = load_workbook(out.getvalue())
    assert time.perf_counter() - t < 10
    assert [b.booking_id for b in res.bookings] == ["BK-1"]
    assert any(f"{MAX_BLANK_RUN:,} empty rows" in n for n in res.notes)
