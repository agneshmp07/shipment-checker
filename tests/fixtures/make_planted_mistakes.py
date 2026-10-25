"""Generate tests/fixtures/planted_mistakes.xlsx in the template's Sheets 1 & 2 format.

EXPECTED is the answer key: the exact set of rule IDs each booking should
break. Run:  python -m tests.fixtures.make_planted_mistakes
"""
from datetime import datetime
from pathlib import Path

import openpyxl

OUT = Path(__file__).with_name("planted_mistakes.xlsx")
# When the case template is present, the fixture is written INTO it (keeping the grey
# example row 3 and Sheets 3-6), exactly as judges will fill it in from row 4.
TEMPLATE = Path(__file__).resolve().parents[2] / "Case2_Shipment_Booking_Test_Data_Format.xlsx"

BOOKING_HEADERS = ["Booking ID", "Seller name", "Origin city", "Origin country", "Destination city",
                   "Destination country", "Destination type", "Ready to ship date", "Seller priority",
                   "Pallets", "Cartons", "Total declared value", "Currency", "Total gross weight (kg)",
                   "Delivery appointment booked"]
ITEM_HEADERS = ["Booking ID", "SKU", "Description", "HS code", "Quantity (units)", "Unit price",
                "Units per carton", "Cartons", "Unit weight (kg)", "Country of origin"]

UK = "United Kingdom"
W, BA = "Warehouse", "Business address"

# (id, seller, origin, dest city, dest country, type, ready, priority, pallets, cartons, value, gross, appt)
BOOKINGS = [
    # Clean bookings
    ("BK-101", "Northfield Goods", "Manchester", "Frankfurt", "Germany", W, "05-12-2026", "Cheapest", 3, 90, 5400, 620, "Yes"),
    ("BK-102", "Thread & Co", "London", "Paris", "France", BA, datetime(2026, 12, 8), "Balanced", 1, 22, 6200, 230, "No"),
    ("BK-103", "Paperworks Ltd", "Birmingham", "Amsterdam", "Netherlands", W, "10-12-2026", "Fastest", 2, 20, 2410, 560, "Yes"),
    # One rule each
    ("BK-104", "Calm Living", "Leeds", "Berlin", "Germany", W, "07-12-2026", "Most reliable", 1, 35, 3200, 320, "Yes"),
    ("BK-105", "Bean There", "London", "Rotterdam", "Netherlands", BA, "09-12-2026", "Cheapest", 1, 15, 2700, 340, "No"),
    ("BK-106", "Kitchen Cloths UK", "Manchester", "Lyon", "France", BA, "11-12-2026", "Balanced", 1, 20, 3500, 130, "No"),
    ("BK-107", "Northfield Goods", "Birmingham", "Hamburg", "Germany", W, "12-12-2026", "Cheapest", 3, 90, 5400, 620, "Yes"),
    ("BK-108", "Glass & Clay", "London", "Amsterdam", "Netherlands", W, "14-12-2026", "Fastest", 3, 60, 6000, 820, "Yes"),
    ("BK-109", "Iron Kitchen", "Leeds", "Frankfurt", "Germany", W, "15-12-2026", "Cheapest", 1, 180, 12960, 1950, "Yes"),
    ("BK-110", "Gadget Gifts", "London", "Paris", "France", BA, "16-12-2026", "Cheapest", 1, 20, 3400, 160, "No"),
    ("BK-111", "Shelf Life Books", "Manchester", "Berlin", "Germany", W, "17-12-2026", "Most reliable", 4, 50, 10500, 900, "No"),
    ("BK-113", "Bright Desk", "Birmingham", "Lyon", "France", BA, "15-10-2026", "Fastest", 1, 25, 2200, 230, "No"),
    # Several errors in one booking
    ("BK-112", "Fresh Start Toiletries", "Leeds", "Rotterdam", "Netherlands", W, "01-09-2026", "Balanced", 1, 43, 2500, 200, "No"),
    # Bad input (not a rulebook mistake): text in the Pallets column
    ("BK-114", "Typo Traders", "London", "Berlin", "Germany", BA, "18-12-2026", "Cheapest", "two", 10, 500, 60, "No"),
]

# (booking, sku, description, hs, qty, unit price, units/carton, cartons, unit kg, country of origin)
ITEMS = [
    ("BK-101", "MUG-01", "Ceramic coffee mug, 350 ml", 691200, 1080, 5.00, 12, 90, 0.5, UK),
    ("BK-102", "TSH-01", "Cotton t-shirt, crew neck", 610910, 600, 8.00, 50, 12, 0.2, "India"),
    ("BK-102", "BAG-02", "Canvas tote bag", 420222, 400, 3.50, 40, 10, 0.15, "India"),
    ("BK-103", "NB-01", "A5 lined notebook", 482010, 2000, 1.20, 100, 20, 0.25, "China"),   # value within 1%
    ("BK-104", "CANDLE-01", "Soy wax candle in glass jar", None, 500, 4.00, 20, 25, 0.3, UK),  # R-01 blank
    ("BK-104", "MAT-02", "Yoga mat, 6 mm", 9506, 100, 12.00, 10, 10, 1.2, "China"),          # R-01 4 digits
    ("BK-105", "COF-01", "Green coffee beans, 1 kg bag", 90111, 300, 9.00, 20, 15, 1.0, "Vietnam"),  # R-01 lost 0
    ("BK-106", "TEA-01", "Cotton tea towel", 630260, 1000, 3.00, 50, 20, 0.1, "Turkey"),     # R-02 3500 vs 3000
    ("BK-107", "MUG-01", "Ceramic coffee mug, 350 ml", 691200, 1080, 5.00, 12, 80, 0.5, UK),  # R-03 + R-04
    ("BK-108", "GLS-01", "Glass flower vase", 701399, 240, 15.00, 6, 40, 1.5, "Germany"),
    ("BK-108", "PLT-02", "Ceramic dinner plate", 691200, 600, 4.00, 24, 25, 0.6, UK),        # R-04 60 vs 65
    ("BK-109", "PAN-01", "Cast iron frying pan", 732391, 720, 18.00, 4, 180, 2.5, "China"),   # R-05 1950 kg/pallet
    ("BK-110", "PB-10", "Portable power bank 10000mAh", 850760, 200, 8.00, 20, 10, 0.3, "China"),  # R-06
    ("BK-110", "SCARF-01", "Wool scarf", 611720, 300, 6.00, 30, 10, 0.2, None),               # R-09
    ("BK-111", "BOOK-01", "Hardback cookery book", 490199, 1500, 7.00, 30, 50, 0.8, UK),     # R-07 900 < 1200; R-08
    ("BK-113", "LAMP-01", "Battery-free LED desk lamp", 940520, 100, 22.00, 4, 25, 1.8, "China"),  # R-10 only
    ("BK-112", "DEO-01", "Deodorant spray 150 ml", 330720, 600, 2.50, 24, 25, 0.2, "Turkey"),  # R-06
    ("BK-112", "SOAP-02", "Handmade soap bar", None, 400, 3.00, 20, 18, 0.12, None),          # R-01, R-03, R-09
    ("BK-114", "CUP-01", "Paper cups, pack of 50", 482369, 100, 5.00, 10, 10, 0.5, UK),
    ("BK-999", "GHOST-1", "Item for a booking that doesn't exist", 691200, 10, 1.00, 10, 1, 0.1, UK),
]

EXPECTED = {
    "BK-101": set(), "BK-102": set(), "BK-103": set(),
    "BK-104": {"R-01"}, "BK-105": {"R-01"}, "BK-106": {"R-02"}, "BK-107": {"R-03", "R-04"},
    "BK-108": {"R-04"}, "BK-109": {"R-05"}, "BK-110": {"R-06", "R-09"}, "BK-111": {"R-07", "R-08"},
    "BK-113": {"R-10"},
    "BK-112": {"R-01", "R-02", "R-03", "R-06", "R-08", "R-09", "R-10"},
    "BK-114": {"DATA"},
}
CLEAN = {"BK-101", "BK-102", "BK-103"}


def _rows():
    bookings = []
    for bk in BOOKINGS:
        bid, seller, oc, dc, dco, typ, ready, prio, pal, ctn, val, gross, appt = bk
        bookings.append([bid, seller, oc, UK, dc, dco, typ, ready, prio, pal, ctn, val, "GBP", gross, appt])
    return bookings, [list(it) for it in ITEMS]


def build(path: Path = OUT) -> Path:
    bookings, items = _rows()
    if TEMPLATE.exists():
        wb = openpyxl.load_workbook(TEMPLATE)
        for sheet, rows in (("Bookings", bookings), ("Items", items)):
            ws = wb[sheet]
            for r, row in enumerate(rows, start=4):          # data starts at row 4
                for c, value in enumerate(row, start=1):
                    ws.cell(r, c, value)
        wb.save(path)
        return path
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bookings"
    ws.append(BOOKING_HEADERS)
    for row in bookings:
        ws.append(row)
    wi = wb.create_sheet("Items")
    wi.append(ITEM_HEADERS)
    for row in items:
        wi.append(row)
    wb.save(path)
    return path


if __name__ == "__main__":
    print("Wrote", build())
