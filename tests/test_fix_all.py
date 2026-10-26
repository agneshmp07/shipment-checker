"""Fix all (high-confidence only) and the corrected Excel download."""
import io
import json
from datetime import date
from pathlib import Path

import openpyxl
from streamlit.testing.v1 import AppTest

from core.check import evaluate
from core.diagnosis import apply_fix, fix_all, high_confidence_actions
from core.excel_export import corrected_workbook
from core.excel_loader import load_workbook
from tests.conftest import STANDIN_CARRIERS, upload_file
from tests.fixtures.make_planted_mistakes import OUT, build
from tests.test_rules import booking, item

TODAY = date(2026, 10, 25)
ROOT = Path(__file__).resolve().parent.parent
ev = lambda b: evaluate(b, STANDIN_CARRIERS, TODAY)


def test_fix_all_only_touches_high_confidence():
    build()
    bookings = load_workbook(OUT).bookings
    changed = fix_all(bookings, ev)
    assert changed == {"BK-106": ["declared value: £3,500.00 → £3,000.00"],
                       "BK-107": ["cartons for MUG-01: 80 → 90"],
                       "BK-108": ["cartons: 60 → 65"]}
    status = {b.booking_id: ev(b).status for b in bookings}
    assert status["BK-106"] == status["BK-107"] == status["BK-108"] == "Ready"
    # left for a person: HS codes, dangerous goods, medium-confidence guesses, slots, dates
    for bid in ("BK-104", "BK-105", "BK-110", "BK-111", "BK-112", "BK-113"):
        assert bid not in changed


def test_total_fix_waits_for_uncertain_item_fix():
    # SOAP-02 style: the item fix is only Medium, so the (dependent) declared-value fix is not applied
    build()
    bk112 = next(b for b in load_workbook(OUT).bookings if b.booking_id == "BK-112")
    assert high_confidence_actions(bk112, ev(bk112).issues) == []


def test_independent_high_fixes_both_apply():
    b = booking([item(cartons=80)], cartons=95)     # item fix High; booking total 95 also wrong
    fix_all([b], ev)
    assert b.items[0].cartons == 90 and b.cartons == 90 and ev(b).status == "Ready"


def test_corrected_file_round_trip_and_highlights():
    build()
    raw = OUT.read_bytes()
    bookings = load_workbook(raw).bookings
    fix_all(bookings, ev)
    bk105 = next(b for b in bookings if b.booking_id == "BK-105")
    apply_fix(bk105, {"target": "item", "sku": "COF-01", "field": "hs_code", "value": "090111"})
    data, changes = corrected_workbook(raw, bookings)
    assert {(c["Booking"], c["Field"], c["Was"], c["Now"]) for c in changes} == {
        ("BK-105", "HS code", "90111", "090111"),
        ("BK-106", "Total declared value", "£3,500.00", "£3,000.00"),
        ("BK-107", "Cartons", "80", "90"),
        ("BK-108", "Cartons", "60", "65"),
    }
    again = {b.booking_id: ev(b).status for b in load_workbook(data).bookings}
    assert again["BK-105"] == again["BK-106"] == again["BK-107"] == again["BK-108"] == "Ready"
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Oops-Proof changes" in wb.sheetnames and "Carrier options" in wb.sheetnames
    hs = next(c for row in wb["Items"].iter_rows() for c in row if c.value == "090111")
    assert hs.fill.start_color.rgb.endswith("FFF2CC") and "90111" in hs.comment.text


def test_no_changes_means_nothing_to_download():
    build()
    raw = OUT.read_bytes()
    data, changes = corrected_workbook(raw, load_workbook(raw).bookings)
    assert changes == [] and "Oops-Proof changes" not in openpyxl.load_workbook(io.BytesIO(data)).sheetnames


def _upload_page(tmp_path, monkeypatch):
    path = tmp_path / "carriers.json"
    path.write_text(json.dumps([c.model_dump() for c in STANDIN_CARRIERS]))
    monkeypatch.setenv("SHIPMENT_CARRIERS_PATH", str(path))
    build()
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60).run()
    at.switch_page("app_pages/upload_file.py").run()
    return upload_file(at, OUT)


def _open_fix_window(at, bid):
    at.session_state["fix_target"] = bid
    return at.run()


def test_fix_window_applies_the_ticked_fixes(tmp_path, monkeypatch):
    at = _upload_page(tmp_path, monkeypatch)
    from ui.correct import suggestions
    bid = next(b.booking_id for b in at.session_state["upload_bookings"]
               if suggestions(evaluate(b, STANDIN_CARRIERS, date.today())))
    _open_fix_window(at, bid)
    assert any(m.value.startswith(f"### {bid}") for m in at.markdown)
    picks = [c for c in at.checkbox if c.key and c.key.startswith(f"pick-{bid}-")]
    assert picks
    for c in picks:
        c.check()
    _open_fix_window(at, bid)                      # the harness re-runs the page, so keep the window open
    at.session_state["fix_target"] = bid
    next(b for b in at.button if b.label.startswith("Apply ")).click().run()
    assert not at.exception
    assert len(at.session_state["fix_all_result"]["changed"][bid]) == len(picks)
    assert any(b.label == "← Back to all problems" for b in at.button)       # still in the fix window


def test_fix_window_change_it_yourself_shows_only_the_fields_needed(tmp_path, monkeypatch):
    at = _upload_page(tmp_path, monkeypatch)
    bk = next(b for b in at.session_state["upload_bookings"]
              if {i.rule_id for i in evaluate(b, STANDIN_CARRIERS, date.today()).issues} == {"R-04"})
    _open_fix_window(at, bk.booking_id)
    labels = {n.label for n in at.number_input}
    assert "Total cartons" in labels                             # R-04 needs only the carton total…
    assert not labels & {"Pallets", "Gross weight (kg)", "Declared value (£)"}   # …not every field
    next(n for n in at.number_input if n.label == "Total cartons").set_value(
        sum(it.cartons for it in bk.items))
    at.session_state["fix_target"] = bk.booking_id
    next(b for b in at.button if b.label == "Save & re-check").click().run()
    assert not at.exception
    assert any("(manual)" in m for m in at.session_state["fix_all_result"]["changed"][bk.booking_id])
    assert any("now passes all 10 checks" in s.value for s in at.success)


def test_fix_window_refuses_values_that_fail_the_checks():
    from ui.fixflow import _invalid
    it = item()
    errors = _invalid({"cartons": 0.0, "total_declared_value": None},
                      [(it, {"hs_code": "123456", "quantity": 2.5, "unit_price": -1.0})])
    text = " ".join(errors)
    assert "isn't a real HS code" in text and "heading 1234" in text
    assert "Cartons must be more than 0" in text and "Declared value is empty" in text
    assert "must be a whole number" in text and "Unit price for" in text
    assert _invalid({"cartons": 80.0}, [(it, {"hs_code": "691200", "quantity": 100.0})]) == []


def test_only_certain_fixes_are_offered(tmp_path, monkeypatch):
    from ui.correct import suggestions
    at = _upload_page(tmp_path, monkeypatch)
    for b in at.session_state["upload_bookings"]:
        for s in suggestions(evaluate(b, STANDIN_CARRIERS, date.today())):
            assert s["conf"] == "High"
            assert s["action"]["field"] not in ("hs_code", "delivery_appointment_booked", "ready_date")
