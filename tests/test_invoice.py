"""Draft commercial invoice: contents, totals, missing fields, PDF and Excel output."""
import io
from datetime import date

import openpyxl
import pytest

from core.check import evaluate
from core.invoice import TODO, InvoiceDetails, build_invoice, invoice_pdf, invoice_xlsx
from tests.conftest import STANDIN_CARRIERS
from tests.test_rules import booking, item

TODAY = date(2026, 10, 25)
FULL = InvoiceDetails(seller_address="12 Mill Lane, Manchester", seller_contact="sales@example.com",
                      seller_vat_eori="GB123456789000", buyer_name="Haus & Herd GmbH",
                      buyer_address="Hanauer Landstr. 1, Frankfurt", buyer_contact="einkauf@example.de",
                      buyer_vat_eori="DE123456789", po_number="PO-778", po_date="20-10-2026",
                      incoterm="DAP", payment_terms="Bank transfer, 30 days")


def inv_for(b, d=FULL):
    return build_invoice(evaluate(b, STANDIN_CARRIERS, TODAY), d, TODAY)


def test_prefilled_from_booking_and_complete():
    inv = inv_for(booking())
    assert inv.number == "INV-BK-001-20261025" and inv.missing == [] and inv.draft_reason == ""
    (line,) = inv.lines
    assert (line.hs_code, line.origin, line.quantity, line.total, line.net_weight_kg) == \
        ("691200", "United Kingdom", 1080, 5400, 540)
    assert ("Marks", "BK-001 / carton 1 to 90") in inv.packages
    assert ("Incoterms", "DAP Frankfurt (Incoterms 2020)") in inv.references


@pytest.mark.parametrize("term,freight", [("DAP", True), ("CPT", True), ("DDP", True), ("EXW", False), ("FCA", False)])
def test_freight_only_when_seller_pays(term, freight):
    inv = inv_for(booking(), FULL.model_copy(update={"incoterm": term}))
    has = any(label.startswith("Freight") for label, _ in inv.charges)
    assert has is freight
    assert inv.invoice_total == (5400 + 450 if freight else 5400)     # stand-in Test C: 3 × £150


def test_discount_and_insurance():
    inv = inv_for(booking(), FULL.model_copy(update={"incoterm": "EXW", "discount_pct": 10, "insurance_gbp": 25}))
    assert inv.charges == [("Discount 10%", -540.0), ("Insurance", 25.0)]
    assert inv.invoice_total == 5400 - 540 + 25


def test_missing_fields_listed_and_marked():
    inv = inv_for(booking(), InvoiceDetails())
    assert "Buyer name" in inv.missing and "Incoterms agreed with the buyer" in inv.missing
    assert "Buyer EU VAT / EORI number" in inv.missing and "Purchase order / proforma date" in inv.missing
    assert inv.buyer[0] == TODO and len(inv.missing) == 11


def test_draft_when_booking_has_red_issues():
    inv = inv_for(booking([item(hs_code=None)]))
    assert inv.draft_reason.startswith("1 issue must be fixed first: HS code missing")
    assert inv.lines[0].hs_code == "MISSING"


def test_pdf_and_excel_output():
    b = booking([item(description="Mug — blue → glaze")])          # characters core PDF fonts lack
    inv = inv_for(b, InvoiceDetails())
    pdf = invoice_pdf(inv)
    assert pdf.startswith(b"%PDF") and len(pdf) > 1500
    wb = openpyxl.load_workbook(io.BytesIO(invoice_xlsx(inv)))
    ws = wb["Commercial invoice"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "COMMERCIAL INVOICE" in values and "691200" in values and 5400 in values
    assert any(isinstance(v, str) and TODO in v for v in values)


def test_po_date_and_buyer_eori_printed():
    inv = inv_for(booking())
    assert ("Purchase order", "PO-778, dated 20-10-2026") in inv.references
    assert "EU VAT / EORI: DE123456789" in inv.buyer


def test_other_charges_line():
    inv = inv_for(booking(), FULL.model_copy(update={"incoterm": "EXW", "other_charges_label": "Chamber certification",
                                                      "other_charges_gbp": 40}))
    assert ("Chamber certification", 40.0) in inv.charges and inv.invoice_total == 5440


OFFICIAL = ("The exporter of the products covered by this document (Exporter Reference Number GB123456789000) "
            "declares that, except where otherwise clearly indicated, these products are of United Kingdom "
            "preferential origin.")


def test_origin_statement_uses_the_official_wording():
    assert inv_for(booking()).origin_statement == []                         # only when the seller ticks it
    inv = inv_for(booking(), FULL.model_copy(update={"origin_statement": True}))
    assert inv.origin_statement[0] == OFFICIAL
    assert inv.origin_statement[-1].endswith("(Name of the exporter)")


def test_origin_statement_excludes_items_not_made_in_the_uk():
    b = booking([item(), item(sku="BOWL-02", country_of_origin="China")], cartons=180, total_declared_value=10800,
                total_gross_weight_kg=1300)
    inv = inv_for(b, FULL.model_copy(update={"origin_statement": True}))
    assert "Not covered by this statement (not of UK origin): BOWL-02." in inv.origin_statement
    only_china = booking([item(country_of_origin="China")])
    assert inv_for(only_china, FULL.model_copy(update={"origin_statement": True})).origin_statement == []


def test_origin_statement_in_pdf_and_excel():
    inv = inv_for(booking(), FULL.model_copy(update={"origin_statement": True}))
    assert invoice_pdf(inv).startswith(b"%PDF")
    ws = openpyxl.load_workbook(io.BytesIO(invoice_xlsx(inv))).active
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert OFFICIAL in texts
