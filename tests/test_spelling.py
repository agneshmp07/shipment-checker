"""Certain spelling corrections: offered and applied only when unambiguous; the issue clears once fixed."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook as open_xlsx

from core.check import evaluate
from core.diagnosis import apply_fix, fix_all
from core.excel_loader import load_workbook
from core.spelling import correct, yes_no
from tests.conftest import STANDIN_CARRIERS
from tests.fixtures.make_planted_mistakes import OUT, build
from core.models import DESTINATIONS, ITEM_ORIGIN_COUNTRIES, PRIORITIES


@pytest.mark.parametrize("raw,allowed,want", [
    ("Germny", list(DESTINATIONS), "Germany"), ("Holland", list(DESTINATIONS), "Netherlands"),
    ("Cheepest", PRIORITIES, "Cheapest"), ("Chnia", ITEM_ORIGIN_COUNTRIES, "China"),
    ("UK", ITEM_ORIGIN_COUNTRIES, "United Kingdom"),
    ("Spain", list(DESTINATIONS), None), ("Japan", ITEM_ORIGIN_COUNTRIES, None),   # real places we don't serve
    ("xx", PRIORITIES, None),
])
def test_correct_only_when_certain(raw, allowed, want):
    assert correct(raw, allowed) == want


def test_yes_no_typos():
    assert yes_no("Yse") is True and yes_no("not yet") is False and yes_no("maybe") is None


def _misspelt_file(tmp_path):
    build()
    wb = open_xlsx(OUT)
    ws = wb.worksheets[0]
    head = {str(c.value).strip().lower(): c.column for row in ws.iter_rows(max_row=4) for c in row if c.value}
    col = next(v for k, v in head.items() if "destination country" in k)
    row = next(r for r in range(4, ws.max_row + 1) if ws.cell(r, col).value == "Germany")
    ws.cell(row, col).value = "Germny"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_misspelt_country_is_fixed_for_certain(tmp_path):
    bookings = load_workbook(_misspelt_file(tmp_path)).bookings
    b = next(b for b in bookings if b.destination_country == "Germny")
    ev = evaluate(b, STANDIN_CARRIERS, date.today())
    issue = next(i for i in ev.issues if i.fix_action and i.fix_action.get("field") == "destination_country")
    assert issue.fix_action["value"] == "Germany" and issue.fix_action["certain"]
    apply_fix(b, issue.fix_action)
    after = evaluate(b, STANDIN_CARRIERS, date.today())
    assert b.destination_country == "Germany"
    assert not [i for i in after.issues if i.field == "destination_country"]      # the data issue is gone


def test_fix_all_corrects_spelling(tmp_path):
    bookings = load_workbook(_misspelt_file(tmp_path)).bookings
    changed = fix_all(bookings, lambda b: evaluate(b, STANDIN_CARRIERS, date.today()))
    assert any("Germny → Germany" in m for msgs in changed.values() for m in msgs)
